"""Builds report/Project_Report.docx (then converted to PDF with LibreOffice)."""
import json
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

R = json.load(open("docs/results.json"))
E = json.load(open("docs/experiments.json"))
B = R["reliability_buckets"]
d = Document()
sec = d.sections[0]; sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
d.styles["Normal"].font.name = "Calibri"; d.styles["Normal"].font.size = Pt(11)
def H(t, l=1): d.add_heading(t, l)
def P(t, bold=False, italic=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = bold; r.italic = italic; return p
def BL(items):
    for i in items: d.add_paragraph(i, style="List Bullet")
def EQ(t):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run(t); r.italic = True; r.font.name = "Cambria Math"
def FIG(path, w, cap):
    d.add_picture(path, width=Inches(w)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = d.add_paragraph(cap); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.runs[0].italic = True
def TABLE(rows, header=True):
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            t.cell(i, j).text = str(v)
            if i == 0 and header: t.cell(i, j).paragraphs[0].runs[0].bold = True
    d.add_paragraph()

def C(text, font, size, bold=False, italic=False, before=0, after=0):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); r.font.name = font; r.font.size = Pt(size); r.bold = bold; r.italic = italic
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font); return p

C("A Project Report", "Arial", 14, bold=True, before=30, after=14)
C("on", "Calibri", 11, italic=True, after=14)
C("Detection of Alphanumeric Handwritten Values using Fuzzy Logic", "Arial", 16, bold=True, after=16)
p = C("carried out as part of the ", "Calibri", 11, italic=True, after=14)
for txt, b in (("Soft Computing Techniques (INT3141)", True), (" Submitted", False)):
    r = p.add_run(txt); r.font.name = "Calibri"; r.font.size = Pt(11); r.italic = True; r.bold = b
C("by", "Calibri", 11, after=14)
C("Abhijay Shrivastava", "Times New Roman", 12, bold=True)
C("2430030333", "Times New Roman", 12, bold=True, after=10)
C("Aarush Dayal", "Times New Roman", 12, bold=True)
C("2430030303", "Times New Roman", 12, bold=True, after=10)
C("Section: Soft Computing D", "Times New Roman", 12, bold=True, after=36)
C("in partial fulfilment for the award of the degree of", "Calibri", 11, italic=True, after=24)
C("Bachelor of Technology", "Calibri", 12, bold=True, after=14)
C("in", "Calibri", 11, after=10)
C("Information Technology", "Calibri", 12, bold=True, after=10)
d.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.CENTER
d.paragraphs[-1].add_run().add_picture("report/assets/logo.png", width=Inches(3.2))
C("Department of Information Technology", "Times New Roman", 12, bold=True, before=14, after=14)
C("MANIPAL UNIVERSITY JAIPUR", "Times New Roman", 12, bold=True)
C("RAJASTHAN, INDIA", "Times New Roman", 12, bold=True, after=18)
C("Oct 2026", "Times New Roman", 12, bold=True)

LINK = "https://github.com/Abhijay09/Softcomputing-project"
p_ = d.add_paragraph("Code repository: "); p_.paragraph_format.page_break_before = True
rid = d.part.relate_to(LINK, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
r_ = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); col = OxmlElement("w:color"); col.set(qn("w:val"), "0563C1")
rpr.append(col); rpr.append(u); r_.append(rpr)
t_ = OxmlElement("w:t"); t_.text = LINK; r_.append(t_); h.append(r_); p_._p.append(h)

H("1. Introduction")
P("Handwritten character recognition is one of the application areas named in the course hand-out. Handwriting is "
  "imprecise: the same letter written twice never looks identical, and different letters (O/0, S/5, I/1) can look "
  "alike. Such imprecision cannot be handled well by crisp yes/no rules, which is exactly the situation fuzzy logic "
  "was designed for (CO2). In this project we build a system that watches a live camera feed (laptop or USB camera), "
  "finds handwritten characters, and recognises each of them with a fuzzy rule base. For every character the system "
  "also reports how reliable the reading is, computed by a second small fuzzy controller.")
P("Course outcomes addressed: INT3141.2 (apply fuzzy logic and reasoning) and INT3141.5 (develop and test a soft "
  "computing technique on a real-life problem).")

H("2. Formal Problem Description")
H("2.1 Input", 2)
P("A video frame is a function I : {0,…,H−1} × {0,…,W−1} → {0,…,255}³ (colour image, H×W pixels; 480×640 for the "
  "laptop camera). The frame contains n ≥ 0 handwritten characters written in dark ink on a light background. "
  "Training data: a labelled set D = {(g_i, y_i)}, where g_i ∈ [0,1]^{28×28} is a grey-scale character image and "
  "y_i ∈ A is its label.")
P("Alphabet: A is the set of 47 EMNIST-balanced classes – the digits 0–9, the upper-case letters A–Z (some letters "
  "that look the same in both cases, e.g. C/c, O/o, S/s, are one class) and the lower-case letters a, b, d, e, f, g, h, n, q, r, t.")
H("2.2 Output", 2)
P("An ordered sequence of n pairs ((c₁, ρ₁), …, (c_n, ρ_n)) where c_k ∈ A is the recognised character at position k "
  "(reading order: left to right, top to bottom) and ρ_k ∈ [0,100] is its fuzzy reliability in percent. The text "
  "output is the string s = c₁c₂…c_n with a space inserted where the horizontal gap between two characters is more "
  "than 1.2 × the median character width.")
H("2.3 Mathematical formulation", 2)
P("Step 1 – Segmentation. The frame is converted to grey, blurred, and thresholded with a local (adaptive) threshold, "
  "giving a binary ink mask M(x,y) ∈ {0,1}. Each 8-connected component of M whose area and height are within limits is "
  "a character candidate with bounding box b_k. Each candidate is resized to fit a 20×20 box (aspect ratio kept) and "
  "centred in a 28×28 image g_k – the same format as the training data.")
P("Step 2 – Feature vector. A feature map φ : [0,1]^{28×28} → [0,1]^{177} gives x = φ(g). The first 49 numbers are the "
  "average ink in each 4×4 block (ink-density zones); the other 128 are stroke-direction values: the image is cut "
  "into 4×4 cells and in each cell the gradient strength is added up in 8 orientation bins (0°–180°), then the whole "
  "vector is scaled to unit length.")
P("Step 3 – Fuzzy rule base. For every class k there are m = 4 fuzzy rules R_{k,r} (188 rules in total):")
EQ("R_{k,r}:  IF x₁ is ~c₁ AND x₂ is ~c₂ AND … AND x₁₇₇ is ~c₁₇₇  THEN class = k")
P("Each antecedent “x_j is ~c_j” is a Gaussian membership function μ_j(x_j) = exp(−(x_j − c_j)² / (2σ_j²)). "
  "The centres c and widths σ of the rules come from fuzzy c-means clustering (Section 3.2). AND is implemented as "
  "the geometric mean of the memberships, so the firing strength of a rule is")
EQ("f_{k,r}(x) = ( Π_j μ_j(x_j) )^{1/177} = exp( −(1/177) Σ_j (x_j − c_j)² / (2σ_j²) )")
P("Rules of the same class are combined with OR = max, and the class membership is μ_k(x) = max_r f_{k,r}(x). "
  "Defuzzification (maximum-membership): ĉ = argmax_k μ_k(x). Strength s = μ_ĉ(x); gap γ = μ_ĉ(x) − (second largest μ_k(x)).")
P("Step 4 – Reliability controller (Mamdani). Inputs s ∈ [0,1] and γ ∈ [0,1] are fuzzified with triangular sets "
  "(Figure 4), nine IF-THEN rules (e.g. IF strength is HIGH AND gap is LARGE THEN reliability is HIGH) are "
  "evaluated with AND = min, clipped outputs are joined with max, and the centroid of the result gives ρ ∈ [0,100].")
H("2.4 Objective and performance measures", 2)
P("Find the rule parameters that maximise top-1 accuracy Acc = (1/N) Σ [ĉ(x_i) = y_i] on a held-out test set, and "
  "make ρ meaningful: samples with higher ρ should be recognised correctly more often. Secondary measure: top-3 "
  "accuracy (true class among the three largest μ_k).")
H("2.5 Assumptions and scope", 2)
BL(["Dark ink on light paper, one or more lines of separate (non-cursive) characters, reasonably large and well lit.",
    "Only the 47 classes above; no punctuation. Touching characters are not split.",
    "Recognition is character-by-character; there is no dictionary or language model."])

H("3. Methodology")
FIG("docs/figs/blocks.png", 6.5, "Figure 1: System block diagram")
H("3.1 Dataset", 2)
P("EMNIST-balanced (Cohen et al., 2017): 112,800 training and 18,800 test images, 28×28, 47 classes with equal "
  "numbers of samples. It is derived from NIST handwriting and is the letters-and-digits extension of MNIST.")
FIG("docs/figs/samples.png", 5.6, "Figure 2: EMNIST samples. Each row is one character class written by different people, so the same character varies a lot in shape and thickness.")
H("3.2 Learning the fuzzy rules with fuzzy c-means", 2)
P("For each class, its training feature vectors are clustered by fuzzy c-means (FCM, fuzzifier 2, 4 clusters). Unlike "
  "k-means, every sample belongs to every cluster with a degree u_{ik} ∈ [0,1], Σ_k u_{ik} = 1. FCM repeats two steps: "
  "(a) c_k = Σ u_{ik}² x_i / Σ u_{ik}²; (b) u_{ik} = 1 / Σ_l (d_{ik}/d_{il})². Each cluster becomes one rule: its centre "
  "gives c_j and the membership-weighted standard deviation gives σ_j. A floor of 0.3 × the overall standard deviation of "
  "feature j is applied, so a rule is never unrealistically narrow. The idea behind several rules per class is that one "
  "letter has several writing styles (e.g. 4 vs. open-top 4, 7 with or without a bar). Section 4.2 measures how much this actually helps.")
FIG("docs/figs/rules_A.png", 5.0, "Figure 3: the four fuzzy rules learned for the letter A (ink-zone part of each rule centre). The four rules turn out to be very similar to each other, which matches the experiment in Section 4.2 where extra rules help only a little.")
H("3.3 Reliability fuzzy controller", 2)
TABLE([["Strength \\ Gap", "SMALL", "MEDIUM", "LARGE"], ["LOW", "LOW", "LOW", "MEDIUM"], ["MEDIUM", "LOW", "MEDIUM", "HIGH"], ["HIGH", "MEDIUM", "HIGH", "HIGH"]])
FIG("docs/figs/fuzzy_sets.png", 6.5, "Figure 4: membership functions of the reliability controller and the distribution of ρ on the test set")
H("3.4 Camera application", 2)
P("A Tkinter desktop application reads frames with OpenCV, runs the pipeline of Figure 1 about 6 times per second, and "
  "draws a box around each character, coloured green/amber/red by reliability level (≥65 / 35–65 / <35 %). The side "
  "panel shows the detected text, the average reliability bar, and for each character the three best classes with their "
  "membership values. Buttons: Freeze/Resume, Save snapshot, Open image…, Switch camera.")
FIG("docs/figs/app_window.png", 6.2, "Figure 5: the Tkinter application reading a still image (Open image… mode). In camera mode the left panel shows the live feed instead. Box colour = reliability level; the panel on the right lists the three best classes and their membership for each character.")
FIG("docs/figs/pipeline.png", 6.5, "Figure 6: pipeline on a sample image (characters taken from the EMNIST test set and placed on a paper-like background)")

H("4. Results (progress till date)")
H("4.1 Accuracy and reliability", 2)
TABLE([["Measure", "Value"], ["Top-1 accuracy (18,800 test images, 47 classes)", f"{R['top1']:.1%}"],
       ["Top-3 accuracy", f"{R['top3']:.1%}"], ["Number of fuzzy rules", R["rules"]], ["Random-guess baseline", "2.1 %"]])
P("Does the reliability value mean anything? Test samples grouped by the reliability the controller assigned:")
TABLE([["Reliability group", "Samples", "Accuracy in group"]] + [[k, v["samples"], f"{v['accuracy']:.1%}"] for k, v in B.items()])
P("Accuracy rises from the LOW to the HIGH group, so ρ is a useful warning: when the fuzzy controller says HIGH the "
  "reading is right about nine times out of ten.")
FIG("docs/figs/confusion.png", 5.2, "Figure 7: confusion matrix")
FIG("docs/figs/per_class.png", 6.3, "Figure 8: accuracy per class")
P("Observations. Most errors are between look-alike shapes: 9/q/g, F/f, 1/I/L, O/0, e/a and 3/B (the most frequent confusions in our test run). This is the inherent imprecision that motivated the fuzzy approach; the top-3 "
  "accuracy shows that the correct character is almost always among the top candidates.")

H("4.2 Experiments: what actually matters?", 2)
P("To understand which design choices matter, we changed one thing at a time and measured test accuracy.")
TABLE([["Rules per class", "Test accuracy"]] + [[k, f"{v:.1%}"] for k, v in E["rules"].items()])
TABLE([["Features used (4 rules per class)", "Test accuracy"]] + [[k, f"{v:.1%}"] for k, v in E["features"].items()])
FIG("docs/figs/experiments.png", 6.3, "Figure 9: effect of the number of rules and of the feature type")
P("Findings. (1) The stroke-direction features carry most of the information: using them alone gives "
  f"{E['features']['stroke directions only (128)']:.1%}, while the ink zones alone give only {E['features']['ink zones only (49)']:.1%}. "
  "Adding the ink zones to the directions improves the result slightly. (2) More rules per class helps only a little "
  f"({E['rules']['1']:.1%} with one rule, {E['rules']['8']:.1%} with eight), so the limit is the feature description and the "
  "simple Gaussian membership form, not the number of writing styles. We kept 4 rules per class as a compromise between "
  "accuracy and the size of the rule base. These findings point to where improvement should come from next (better "
  "features, tuned membership widths).")
H("5. Limitations and future work")
BL(["Accuracy (about 66 %) is lower than a trained neural network would reach (CO3). The fuzzy rules are deliberately simple and "
    "explainable: every decision can be traced back to rules and membership values.",
    "Future: tune the rule widths and number of rules per class with a genetic algorithm (CO4); add a perceptron / neural "
    "network and compare; use context (dictionary) to resolve O/0 and I/1; split touching characters.",
    "Camera robustness: strong shadows, thin pen strokes and tilted paper reduce quality."])

H("6. Tools and how to run")
P("Python 3, NumPy (fuzzy c-means and fuzzy rules written from scratch), OpenCV (image processing), Tkinter + Pillow (GUI), "
  "Matplotlib (figures). Commands:  python src/train.py  |  python src/evaluate.py  |  python src/app.py [camera index or image file]")

H("7. References")
BL(["L. A. Zadeh, “Fuzzy sets”, Information and Control, 8(3), 1965.",
    "J. C. Bezdek, Pattern Recognition with Fuzzy Objective Function Algorithms, Plenum, 1981 (fuzzy c-means).",
    "E. H. Mamdani and S. Assilian, “An experiment in linguistic synthesis with a fuzzy logic controller”, 1975.",
    "G. Cohen, S. Afshar, J. Tapson, A. van Schaik, “EMNIST: an extension of MNIST to handwritten letters”, 2017.",
    "S. Rajasekaran and G. A. Vijayalakshmi Pai, Neural Networks, Fuzzy Logic and Genetic Algorithms, PHI.",
    "N. Dalal and B. Triggs, “Histograms of oriented gradients for human detection”, CVPR 2005 (idea behind direction features)."])
d.save("report/Project_Report.docx")
